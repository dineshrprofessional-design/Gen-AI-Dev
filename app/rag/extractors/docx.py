"""DOCX extraction.

Better than PDF: paragraphs and table cells are real objects, so headings keep
their level and table rows keep their shape.

The body is walked in **document order**. `document.paragraphs` and
`document.tables` are two separate lists, so reading them one after the other
puts every table at the end of the text — detached from the heading it belongs
under. That would orphan a parameter table at ingest time, before any chunker
got a chance to keep it together.
"""

from pathlib import Path

from app.rag.extractors.base import Extracted, ExtractError


def _iter_body(document):
    """Yield paragraphs and tables in the order they appear in the document."""
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            yield Paragraph(child, document)
        elif child.tag.endswith("}tbl"):
            yield Table(child, document)


def _heading_prefix(style_name: str) -> str:
    """Map a Word paragraph style to a markdown heading prefix, or ''."""
    if style_name == "Title":
        return "# "
    if style_name.startswith("Heading"):
        tail = style_name.rsplit(" ", 1)[-1]
        level = int(tail) if tail.isdigit() else 1
        return "#" * min(level, 6) + " "
    return ""


def _render_table(table) -> str:
    """Render a table as contiguous markdown lines.

    Contiguous matters: a blank line between rows ends the table as far as any
    markdown parser is concerned, and the rows stop being a table at all.
    """
    lines: list[str] = []
    for index, row in enumerate(table.rows):
        cells = [c.text.strip() for c in row.cells]
        if not any(cells):
            continue
        lines.append("| " + " | ".join(cells) + " |")
        if index == 0:
            lines.append("| " + " | ".join("---" for _ in cells) + " |")
    return "\n".join(lines)


class DocxExtractor:
    name = "docx"
    extensions = (".docx",)

    def extract(self, path: Path) -> Extracted:
        try:
            import docx
        except ImportError as exc:  # pragma: no cover - environment dependent
            raise ExtractError("python-docx is not installed") from exc

        try:
            document = docx.Document(str(path))
        except Exception as exc:
            raise ExtractError(f"unreadable DOCX: {exc}") from exc

        parts: list[str] = []
        title_hint: str | None = None

        for item in _iter_body(document):
            if hasattr(item, "rows"):  # a table
                rendered = _render_table(item)
                if rendered:
                    parts.append(rendered)
                continue

            text = item.text.strip()
            if not text:
                continue
            style = item.style.name if item.style else ""
            prefix = _heading_prefix(style)
            if prefix and title_hint is None:
                title_hint = text
            parts.append(prefix + text)

        body = "\n\n".join(parts)
        if not body.strip():
            raise ExtractError("DOCX contained no text")

        return Extracted(
            text=body,
            title_hint=title_hint or (parts[0] if parts else None),
        )
